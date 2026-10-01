#!/usr/bin/env python3
"""Convierte un Documento de Entrada de Datos (DED) de gobierno al contrato ODCS v3.2.0.

Entrada:
  - El Excel DED tal como lo usa gobierno (hojas "Informacion - tabla" y "DED").
  - Opcional: un complemento YAML con lo que el DED no cubre (versión del contrato,
    tenant, dataProduct, servidores, reglas de calidad de negocio). Se busca junto al Excel:
    <nombre>.contrato.yaml

Salida:
  - contracts/<flujo>/<nombre_tabla_landing>.odcs.yaml

Principio: falla de forma ruidosa. Si el DED no está completo o tiene valores fuera de
las listas, el conversor se detiene y lista TODOS los errores con hoja, fila y el rol
responsable de corregirlos. No adivina ni completa datos.

Uso:
  python scripts/ded_a_contrato.py alcance/landing_ded/de_maestro_canal.xlsx
  python scripts/ded_a_contrato.py <ded.xlsx> --output <contrato.odcs.yaml>
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
import warnings
from pathlib import Path

import openpyxl
import yaml

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

HOJA_TABLA = "Informacion - tabla"
HOJA_CAMPOS = "DED"
HOJA_VERSIONES = "Historial de versiones"
VERSIONES_PLANTILLA_SOPORTADAS = {"3"}

# ── Encabezados del DED (normalizados: sin tildes, minúsculas, sin ¿?) ──────────────
COLS_TABLA = {
    "ruta_broad": "ruta en broad",
    "contenedor_broad": "nombre del contenedor broad",
    "dominio": "dominio / subdominio",
    "capa": "capa",
    "empresa": "empresa",
    "schema_landing": "schema landing/ archivo",
    "tabla_fuente": "nombre tabla fuente",
    "tabla_landing": "nombre de tabla/ archivo landing",
    "registros": "cantidad de registros",
    "historia": "historia",
    "ruta_landing": "ruta landing",
    "descripcion": "descripcion de la tabla landing",
    "tipo_fuente": "tipo de fuente origen",
    "tipo_carga": "tipo de carga",
    "frecuencia": "frecuencia de actualizacion landing",
    "dac": "dac",
    "product_owner": "product owner",
    "squad": "squad",
}
COLS_CAMPOS = {
    "id": "id",
    "fuente": "fuente",
    "trazabilidad": "aplica trazabilidad",
    "nombre_fuente": "nombre campo",
    "nombre_modelado": "nombre campo modelado",
    "tipo_formato": "tipo de formato",
    "formato_input": "formato input",
    "formato_landing": "formato landing",
    "definicion": "definicion del campo",
    "longitud_fuente": "longitud del campo fuente",
    "longitud_landing": "longidud del campo landing",  # así está escrito en la plantilla
    "ejemplos": "ejemplo de valores",
    "catalogo": "catalogo dominio de valores",
    "descartado": "campo descartado",
    "pk": "es una llave primaria",
    "llave_tecnica": "es una llave tecnica",
    "dac": "es un dato dac",
    "vacios": "el campo contiene vacios",
    "nulos": "el campo contiene nulos",
    "criticidad": "criticidad",
    "sustento_criticidad": "sustento de criticidad",
    "clasificacion": "nivel de clasificacion del dato",
}
ALIAS_ENCABEZADO = {"longitud del campo landing": "longidud del campo landing"}

# Campos que el conversor exige aunque la plantilla marque todo como OBLIGATORIO
# (en la práctica "Catálogo" o "Longitud" suelen venir vacíos o con NO APLICA).
OBLIGATORIOS_TABLA = ["dominio", "capa", "empresa", "schema_landing", "tabla_fuente",
                      "tabla_landing", "descripcion", "tipo_fuente", "tipo_carga", "frecuencia"]
OBLIGATORIOS_CAMPO = ["nombre_fuente", "nombre_modelado", "formato_landing", "definicion",
                      "descartado", "pk", "dac", "nulos", "clasificacion"]

# ── Listas de valores (hoja "Documento - valores") ──────────────────────────────────
DOMINIOS = {"Atenciones Médicas", "Canales de Venta", "Cliente Cuenta Persona Jurídica",
            "Cliente Cuenta Persona Natural", "Cobranzas", "Data & Analytics",
            "Empresas Prestadoras", "Finanzas", "Médicos", "Póliza", "Producto Empresa",
            "Producto Persona", "Siniestros"}
EMPRESAS = {"demo", "corp"}  # valores demo; los reales no se publican
SCHEMAS_LANDING = {"dataentry_int": "transitorio (ingesta interna)",
                   "dataentry_ext": "información externa",
                   "dataentry_usr": "permanente (usuario)"}
TIPOS_FUENTE = {"DESCRIPTIVA", "MAESTRA", "HISTORICA", "DATA ENTRY"}
TIPOS_CARGA = {"Full", "Incremental"}
FRECUENCIAS = {"DIARIA (L - V) Feriados", "DIARIA (L - V) Sin Feriados",
               "DIARIA (L - D) Feriados", "DIARIA (L - D) Sin Feriados",
               "SEMANAL", "MENSUAL", "A DEMANDA"}
CLASIFICACIONES = {"publico": "public", "uso interno": "internal", "restringido": "restricted"}
CRITICIDADES = {"EDC", "EDNC"}

# Formato Landing del DED -> (logicalType ODCS, tipo físico Databricks)
TIPOS = {
    "string": ("string", "STRING"),
    "varchar": ("string", "STRING"),
    "char": ("string", "STRING"),
    "decimal": ("number", None),  # DECIMAL(p,s) sale de "Longitud del campo Landing"
    "date": ("date", "DATE"),
    "time_stamp": ("timestamp", "TIMESTAMP"),
    "timestamp": ("timestamp", "TIMESTAMP"),  # alias usado en la práctica
    "boolean": ("boolean", "BOOLEAN"),
    "bit": ("boolean", "BOOLEAN"),
    "time": ("string", "STRING"),  # Databricks no tiene tipo TIME
}
NOMBRE_VALIDO = re.compile(r"^[a-z][a-z0-9_]*$")


class ErroresDED(Exception):
    pass


class _Dumper(yaml.SafeDumper):
    """Pone comillas a textos que parecen números o booleanos (ej. '0000029', 'NO')."""


def _repr_str(dumper, data):
    ambiguo = re.fullmatch(r"[-+]?[0-9_.,:eE]+|yes|no|on|off|true|false|null|~", data, re.IGNORECASE)
    return dumper.represent_scalar("tag:yaml.org,2002:str", data, style="'" if ambiguo else None)


_Dumper.add_representer(str, _repr_str)


def norm(texto) -> str:
    """Normaliza un encabezado: sin tildes, minúsculas, sin ¿?, espacios simples."""
    if texto is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = t.replace("¿", "").replace("?", "").replace("\n", " ").lower()
    return re.sub(r"\s+", " ", t).strip()


def texto(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def si_no(v) -> str | None:
    """'SÍ', 'Si', 'si' -> 'SI'; 'no' -> 'NO'; otro -> None."""
    n = norm(v).upper()
    return n if n in {"SI", "NO"} else None


def es_no_aplica(v) -> bool:
    return norm(v) in {"", "no aplica", "n/a", "na", "-"}


class Lector:
    def __init__(self, ruta: Path):
        self.ruta = ruta
        self.errores: list[str] = []
        self.avisos: list[str] = []
        self.etiquetas: dict[str, str] = {}
        self.wb = openpyxl.load_workbook(ruta)  # fórmulas
        self.wv = openpyxl.load_workbook(ruta, data_only=True)  # valores calculados por Excel

    def error(self, msg):
        self.errores.append(msg)

    def aviso(self, msg):
        self.avisos.append(msg)

    def hoja(self, nombre):
        if nombre not in self.wb.sheetnames:
            raise ErroresDED(f"No existe la hoja '{nombre}'. ¿Es un Documento de Entrada de Datos (DED)?")
        return self.wb[nombre], self.wv[nombre]

    def valor(self, hoja_f, hoja_v, fila, col):
        """Valor de una celda. Si es fórmula usa el valor calculado; si no lo hay, lo resuelve."""
        crudo = hoja_f.cell(fila, col).value
        if isinstance(crudo, str) and crudo.startswith("="):
            cache = hoja_v.cell(fila, col).value
            return cache if cache is not None else ("__FORMULA__", crudo)
        return crudo

    def mapear_encabezados(self, ws, fila, esperadas, hoja_nombre):
        hallados = {}
        for c in range(1, ws.max_column + 1):
            h = norm(ws.cell(fila, c).value)
            h = ALIAS_ENCABEZADO.get(h, h)
            for clave, esperado in esperadas.items():
                if clave not in hallados and (h == esperado or (h.startswith(esperado) and esperado.startswith("tipo de fuente"))):
                    hallados[clave] = c
        self.etiquetas.update({k: re.sub(r"\s+", " ", str(ws.cell(fila, c).value)).strip() for k, c in hallados.items()})
        faltan = [esperadas[k] for k in esperadas if k not in hallados]
        if faltan:
            raise ErroresDED(f"Hoja '{hoja_nombre}' fila {fila}: no se encontraron las columnas {faltan}. "
                             "¿Cambió la plantilla?")
        return hallados

    def buscar_fila(self, ws, encabezado_norm):
        for r in range(1, min(ws.max_row, 40) + 1):
            for c in range(1, ws.max_column + 1):
                if norm(ws.cell(r, c).value) == encabezado_norm:
                    return r
        return None


def version_plantilla(lec: Lector) -> str | None:
    if HOJA_VERSIONES not in lec.wb.sheetnames:
        lec.aviso("No se encontró la hoja 'Historial de versiones'.")
        return None
    ws = lec.wb[HOJA_VERSIONES]
    versiones = []
    for row in ws.iter_rows():
        for c in row:
            if norm(c.value) == "version":
                r = c.row + 1
                while ws.cell(r, c.column).value not in (None, ""):
                    versiones.append(str(ws.cell(r, c.column).value).strip())
                    r += 1
    if not versiones:
        return None
    v = versiones[-1]
    if v not in VERSIONES_PLANTILLA_SOPORTADAS:
        lec.aviso(f"Versión de plantilla {v} no probada con este conversor "
                  f"(soportadas: {sorted(VERSIONES_PLANTILLA_SOPORTADAS)}).")
    return v


def leer_tabla(lec: Lector) -> dict:
    wf, wv = lec.hoja(HOJA_TABLA)
    fila_h = lec.buscar_fila(wf, COLS_TABLA["tabla_landing"])
    if not fila_h:
        raise ErroresDED(f"Hoja '{HOJA_TABLA}': no se encontró el encabezado 'NOMBRE DE TABLA/ ARCHIVO LANDING'.")
    cols = lec.mapear_encabezados(wf, fila_h, COLS_TABLA, HOJA_TABLA)
    fila_roles, fila_d = fila_h - 2, fila_h + 1
    rol = {k: texto(wf.cell(fila_roles, c).value) or "?" for k, c in cols.items()}
    t = {k: lec.valor(wf, wv, fila_d, c) for k, c in cols.items()}
    t = {k: (None if isinstance(v, tuple) else texto(v)) for k, v in t.items()}
    donde = f"'{HOJA_TABLA}' fila {fila_d}"

    for k in OBLIGATORIOS_TABLA:
        if not t[k]:
            lec.error(f"{donde}: falta '{lec.etiquetas.get(k, COLS_TABLA[k])}' (responsable: {rol[k]})")
    if t["tabla_landing"] and not NOMBRE_VALIDO.match(t["tabla_landing"]):
        lec.error(f"{donde}: el nombre de tabla landing '{t['tabla_landing']}' debe ser snake_case en minúsculas "
                  f"(responsable: {rol['tabla_landing']})")
    reglas = [("dominio", DOMINIOS), ("empresa", EMPRESAS), ("schema_landing", set(SCHEMAS_LANDING)),
              ("tipo_fuente", TIPOS_FUENTE), ("tipo_carga", TIPOS_CARGA), ("frecuencia", FRECUENCIAS)]
    for k, permitidos in reglas:
        if t[k] and t[k] not in permitidos:
            lec.error(f"{donde}: '{lec.etiquetas.get(k, COLS_TABLA[k])}' = '{t[k]}' no está en la lista de valores "
                      f"(responsable: {rol[k]})")
    if t["capa"] and t["capa"].upper() != "LANDING":
        lec.error(f"{donde}: la capa debe ser LANDING para un DED (valor: '{t['capa']}')")
    t["_fila"] = fila_d
    return t


def leer_campos(lec: Lector) -> list[dict]:
    wf, wv = lec.hoja(HOJA_CAMPOS)
    fila_h = lec.buscar_fila(wf, COLS_CAMPOS["nombre_modelado"])
    if not fila_h:
        raise ErroresDED(f"Hoja '{HOJA_CAMPOS}': no se encontró el encabezado 'Nombre Campo Modelado'.")
    cols = lec.mapear_encabezados(wf, fila_h, COLS_CAMPOS, HOJA_CAMPOS)
    fila_roles = next((r for r in range(1, fila_h) if norm(wf.cell(r, 2).value) == "responsable"), None)
    rol = {k: (texto(wf.cell(fila_roles, c).value) if fila_roles else None) or "?" for k, c in cols.items()}

    campos = []
    r = fila_h + 1
    while True:
        nombre_fuente = texto(wf.cell(r, cols["nombre_fuente"]).value)
        fila_vacia = all(wf.cell(r, c).value in (None, "") or
                         (isinstance(wf.cell(r, c).value, str) and wf.cell(r, c).value.startswith("=(ROW"))
                         for c in cols.values())
        if fila_vacia:
            break
        f = {}
        for k, c in cols.items():
            v = lec.valor(wf, wv, r, c)
            if isinstance(v, tuple):  # fórmula sin valor calculado (archivo no guardado por Excel)
                formula = v[1].upper()
                if k == "nombre_modelado" and formula.startswith("=LOWER("):
                    v = nombre_fuente.lower() if nombre_fuente else None
                elif k == "id":
                    v = r - fila_h
                else:
                    lec.error(f"'{HOJA_CAMPOS}' fila {r}: no se puede resolver la fórmula {v[1]} en "
                              f"'{lec.etiquetas.get(k, COLS_CAMPOS[k])}'. Abre y guarda el archivo en Excel.")
                    v = None
            f[k] = texto(v)
        f["_fila"] = r
        campos.append(f)
        r += 1

    if not campos:
        lec.error(f"'{HOJA_CAMPOS}': el documento no tiene campos.")
    return campos, rol


def validar_campos(lec: Lector, campos: list[dict], rol: dict) -> list[dict]:
    vistos = {}
    activos = []
    for f in campos:
        donde = f"'{HOJA_CAMPOS}' fila {f['_fila']} ({f.get('nombre_fuente') or 'sin nombre'})"
        for k in ["descartado", "pk", "dac", "nulos", "vacios", "llave_tecnica", "trazabilidad"]:
            if f.get(k) is not None:
                n = si_no(f[k])
                if n is None:
                    lec.error(f"{donde}: '{lec.etiquetas.get(k, COLS_CAMPOS[k])}' debe ser SI o NO (valor: '{f[k]}', responsable: {rol[k]})")
                f[k] = n
        if f.get("descartado") == "SI":
            continue
        for k in OBLIGATORIOS_CAMPO:
            if not f.get(k):
                lec.error(f"{donde}: falta '{lec.etiquetas.get(k, COLS_CAMPOS[k])}' (responsable: {rol[k]})")
        nm = f.get("nombre_modelado")
        if nm:
            if not NOMBRE_VALIDO.match(nm):
                lec.error(f"{donde}: 'Nombre Campo Modelado' = '{nm}' debe ser snake_case en minúsculas, "
                          f"sin espacios ni tildes (responsable: {rol['nombre_modelado']})")
            elif nm in vistos:
                lec.error(f"{donde}: 'Nombre Campo Modelado' = '{nm}' está repetido (también en la fila {vistos[nm]})")
            vistos.setdefault(nm, f["_fila"])
        fl = norm(f.get("formato_landing")).replace(" ", "_")
        if f.get("formato_landing"):
            if fl not in TIPOS:
                lec.error(f"{donde}: 'Formato Landing' = '{f['formato_landing']}' no es un tipo válido "
                          f"(usa: Char, Varchar, String, Decimal, Date, Time_Stamp, Boolean) "
                          f"(responsable: {rol['formato_landing']})")
            elif fl == "decimal":
                m = re.fullmatch(r"\s*(\d+)\s*[,.;]\s*(\d+)\s*", f.get("longitud_landing") or "")
                if not m:
                    lec.error(f"{donde}: un campo Decimal requiere 'Longitud del campo Landing' como "
                              f"'enteros,decimales' (ej. 12,2). Valor: '{f.get('longitud_landing')}'")
        if f.get("clasificacion") and norm(f["clasificacion"]) not in CLASIFICACIONES:
            lec.error(f"{donde}: 'Nivel de clasificación' = '{f['clasificacion']}' no está en la lista "
                      f"(Público, Uso interno, Restringido)")
        if f.get("criticidad") and f["criticidad"] not in CRITICIDADES:
            lec.error(f"{donde}: 'Criticidad' = '{f['criticidad']}' debe ser EDC o EDNC")
        if f.get("dac") == "SI" and norm(f.get("clasificacion")) == "publico":
            lec.aviso(f"{donde}: está marcado como DAC pero su clasificación es 'Público'.")
        activos.append(f)
    if activos and not any(f.get("pk") == "SI" for f in activos):
        lec.aviso("Ningún campo está marcado como llave primaria: no se validarán duplicados.")
    return activos


def campo_a_propiedad(f: dict, pos_pk: int | None) -> dict:
    fl = norm(f["formato_landing"]).replace(" ", "_")
    logico, fisico = TIPOS[fl]
    opciones = {}
    if fl == "decimal":
        p, s = re.findall(r"\d+", f["longitud_landing"])[:2]
        fisico = f"DECIMAL({p},{s})"
    elif logico == "string" and f.get("longitud_landing") and f["longitud_landing"].isdigit():
        opciones["maxLength"] = int(f["longitud_landing"])
    clasif = CLASIFICACIONES[norm(f["clasificacion"])]
    if f.get("dac") == "SI" and clasif == "public":
        clasif = "restricted"

    p = {"id": f["nombre_modelado"], "name": f["nombre_modelado"],
         "businessName": f["nombre_fuente"], "logicalType": logico, "physicalType": fisico,
         "description": f["definicion"], "required": f.get("nulos") == "NO", "classification": clasif}
    if pos_pk:
        p["primaryKey"] = True
        p["primaryKeyPosition"] = pos_pk
    if opciones:
        p["logicalTypeOptions"] = opciones
    if f.get("dac") == "SI":
        p["tags"] = ["dac"]
    if f.get("ejemplos") and not es_no_aplica(f["ejemplos"]):
        ej = [e.strip() for e in f["ejemplos"].split(";") if e.strip()]
        p["examples"] = list(dict.fromkeys(ej))[:4]
    if f.get("catalogo") and not es_no_aplica(f["catalogo"]):
        valores = [v.strip() for v in re.split(r"[;,\n]", f["catalogo"]) if v.strip()]
        p["quality"] = [{"type": "library", "metric": "invalidValues", "arguments": {"validValues": valores},
                         "mustBe": 0, "name": f"{f['nombre_modelado']}_catalogo",
                         "description": "Valores permitidos según el catálogo del DED.",
                         "dimension": "conformity", "severity": "error"}]
    props = [("fuente", "fuente"), ("campoTecnico", None), ("dac", "dac"), ("contieneVacios", "vacios"),
             ("criticidad", "criticidad"), ("sustentoCriticidad", "sustento_criticidad"),
             ("tipoFormatoGobierno", "tipo_formato"), ("formatoInput", "formato_input"),
             ("longitudFuente", "longitud_fuente"), ("aplicaTrazabilidad", "trazabilidad"),
             ("llaveTecnica", "llave_tecnica")]
    cp = []
    for nombre, k in props:
        if nombre == "campoTecnico":
            if norm(f.get("fuente")) == "(campo tecnico)":
                cp.append({"property": "campoTecnico", "value": True})
            continue
        v = f.get(k)
        if v is not None and not es_no_aplica(v):
            cp.append({"property": nombre, "value": v})
    if cp:
        p["customProperties"] = cp
    return p


def construir_contrato(t: dict, activos: list[dict], todos: list[dict], version_pl: str | None,
                       complemento: dict, ruta_ded: Path) -> dict:
    tabla = t["tabla_landing"]
    schema_landing = t["schema_landing"]
    props, pos = [], 0
    for f in activos:
        if f.get("pk") == "SI":
            pos += 1
            props.append(campo_a_propiedad(f, pos))
        else:
            props.append(campo_a_propiedad(f, None))
    hay_dac = any(f.get("dac") == "SI" for f in activos)
    descartados = [f["nombre_fuente"] for f in todos if f.get("descartado") == "SI"]

    calidad_tabla = [{"type": "library", "metric": "rowCount", "mustBeGreaterThan": 0,
                      "name": "tabla_no_vacia", "description": "La carga del data entry no puede venir vacía.",
                      "dimension": "completeness", "severity": "error"}]
    calidad_tabla += complemento.get("quality", [])

    modelo = {"id": tabla, "name": tabla, "physicalName": tabla, "physicalType": "table",
              "businessName": t["tabla_fuente"], "description": t["descripcion"],
              "tags": ["dataentry", "dac"] if hay_dac else ["dataentry"],
              "properties": props, "quality": calidad_tabla}

    cp = [("capa", "LANDING"), ("flujo", "DED"), ("empresa", t["empresa"]),
          ("schemaLanding", schema_landing), ("tipoDataEntry", SCHEMAS_LANDING.get(schema_landing)),
          ("nombreTablaFuente", t["tabla_fuente"]), ("tipoFuente", t["tipo_fuente"]),
          ("tipoCarga", t["tipo_carga"]), ("frecuenciaActualizacion", t["frecuencia"]),
          ("historia", t["historia"]), ("registrosCargaInicial", t["registros"]),
          ("rutaBroad", t["ruta_broad"]), ("contenedorBroad", t["contenedor_broad"]),
          ("rutaLanding", t["ruta_landing"]), ("dac", "SI" if hay_dac else "NO"),
          ("productOwner", t["product_owner"]), ("squad", t["squad"]),
          ("camposDescartados", ", ".join(descartados) if descartados else None),
          ("documentoAlcance", ruta_ded.as_posix()), ("versionPlantillaDED", version_pl)]
    custom = [{"property": k, "value": v} for k, v in cp if v not in (None, "") and not es_no_aplica(v)]

    servidores = complemento.get("servers", [])
    for s in servidores:  # el esquema landing lo manda el DED, no el complemento
        if s.get("type") == "databricks":
            s["schema"] = schema_landing

    contrato = {
        "apiVersion": "v3.2.0",
        "kind": "DataContract",
        "id": f"landing_{schema_landing}_{tabla}",
        "name": t["tabla_fuente"],
        "version": str(complemento.get("version", "1.0.0")),
        "status": complemento.get("status", "draft"),
        **{k: complemento[k] for k in ("tenant", "dataProduct") if complemento.get(k)},
        "domain": t["dominio"],
        "tags": ["landing", "ded"],
        "description": {"purpose": t["descripcion"],
                        "usage": f"Data entry {SCHEMAS_LANDING.get(schema_landing, '')} en la capa landing."},
        "customProperties": custom,
    }
    if servidores:
        contrato["servers"] = servidores
    contrato["schema"] = [modelo]
    if t.get("product_owner") or t.get("squad"):
        contrato["team"] = {"name": t.get("squad") or "Equipo de datos",
                            "members": [{"username": t["product_owner"], "role": "product-owner"}]
                            if t.get("product_owner") else []}
    if complemento.get("slaProperties"):
        contrato["slaProperties"] = complemento["slaProperties"]
    return contrato


def convertir(ruta_ded: Path, ruta_salida: Path | None, raiz: Path) -> Path:
    lec = Lector(ruta_ded)
    version_pl = version_plantilla(lec)
    t = leer_tabla(lec)
    campos, rol = leer_campos(lec)
    activos = validar_campos(lec, campos, rol)

    for a in lec.avisos:
        print(f"  ⚠️  {a}")
    if lec.errores:
        print(f"\n❌ El DED '{ruta_ded.name}' tiene {len(lec.errores)} error(es). Corrígelos y vuelve a ejecutar:\n")
        for e in lec.errores:
            print(f"   • {e}")
        raise ErroresDED(f"{len(lec.errores)} error(es) en el DED")

    ruta_comp = ruta_ded.with_suffix(".contrato.yaml")
    complemento = yaml.safe_load(ruta_comp.read_text(encoding="utf-8")) if ruta_comp.exists() else {}

    try:
        rel = ruta_ded.resolve().relative_to(raiz.resolve())
    except ValueError:
        rel = Path(ruta_ded.name)
    contrato = construir_contrato(t, activos, campos, version_pl, complemento or {}, rel)

    if ruta_salida is None:
        ruta_salida = raiz / "contracts" / ruta_ded.parent.name / f"{t['tabla_landing']}.odcs.yaml"
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    cabecera = (f"# GENERADO desde {rel.as_posix()} con scripts/ded_a_contrato.py — NO EDITAR A MANO.\n"
                f"# Para cambiarlo, modifica el DED (o su complemento .contrato.yaml) y regenera.\n")
    with open(ruta_salida, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(cabecera)
        yaml.dump(contrato, fh, Dumper=_Dumper, allow_unicode=True, sort_keys=False, width=110)
    n_desc = sum(1 for c in campos if c.get("descartado") == "SI")
    try:
        mostrar = ruta_salida.resolve().relative_to(raiz.resolve()).as_posix()
    except ValueError:
        mostrar = ruta_salida.as_posix()
    print(f"  ✅ {len(activos)} campos → {mostrar}"
          + (f" ({n_desc} descartado(s) excluido(s))" if n_desc else ""))
    return ruta_salida


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ded", type=Path, help="Ruta al Excel DED")
    ap.add_argument("--output", type=Path, default=None, help="Ruta del contrato de salida")
    ap.add_argument("--root", type=Path, default=Path(__file__).resolve().parent.parent)
    a = ap.parse_args()
    try:
        convertir(a.ded, a.output, a.root)
    except ErroresDED as e:
        if not str(e).endswith("en el DED"):
            print(f"\n❌ {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
